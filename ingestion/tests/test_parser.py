from pathlib import Path

import fitz
import pytest
from docx import Document

from pipelines.parser import (
    UnsupportedDocumentType,
    clean_text,
    extract_text_from_docx,
    extract_text_from_pdf,
    parse_document,
)


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Second trimester nutrition guidance.")
    page.insert_text((72, 96), "Eat iron-rich foods such as spinach and moong dal.")
    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Second trimester nutrition guidance.")
    document.add_paragraph("Eat iron-rich foods such as spinach and moong dal.")
    document.save(path)
    return path


def test_extract_text_from_pdf(sample_pdf: Path) -> None:
    text = extract_text_from_pdf(sample_pdf)
    assert "Second trimester nutrition guidance." in text
    assert "spinach" in text


def test_extract_text_from_docx(sample_docx: Path) -> None:
    text = extract_text_from_docx(sample_docx)
    assert "Second trimester nutrition guidance." in text
    assert "moong dal" in text


def test_parse_document_dispatches_by_extension(sample_pdf: Path, sample_docx: Path) -> None:
    assert "nutrition" in parse_document(sample_pdf)
    assert "nutrition" in parse_document(sample_docx)


def test_parse_document_rejects_unsupported_type(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text("hello")
    with pytest.raises(UnsupportedDocumentType):
        parse_document(path)


def test_clean_text_rejoins_hyphenated_line_break() -> None:
    assert clean_text("This is an exam-\nple sentence.") == "This is an example sentence."


def test_clean_text_collapses_horizontal_whitespace() -> None:
    assert clean_text("Too    many   spaces\there") == "Too many spaces here"


def test_clean_text_collapses_excess_blank_lines_but_keeps_paragraphs() -> None:
    raw = "Paragraph one.\n\n\n\n\nParagraph two."
    assert clean_text(raw) == "Paragraph one.\n\nParagraph two."


def test_clean_text_strips_trailing_whitespace_per_line() -> None:
    raw = "line one   \nline two\t\n"
    assert clean_text(raw) == "line one\nline two"


def test_clean_text_does_not_alter_meaning() -> None:
    raw = "  Dietary advice: avoid raw papaya during pregnancy.  "
    assert clean_text(raw) == "Dietary advice: avoid raw papaya during pregnancy."
